#!/usr/bin/env python3
"""Build the narrated architecture prelude: deck, previews, per-slide voice-over and video.

Source: ``docs/Decks/architecture-prelude/voiceover-script.md`` (format described at its top).
Outputs:

- ``docs/Decks/architecture-prelude/ask-the-tower-architecture-prelude.pptx`` — 16:9, one visual per slide, speaker
  notes = narration, each slide's MP3 embedded and started on slide start, auto-advance after it;
- ``docs/Decks/architecture-prelude/preview/slide-NN.png`` — every slide rendered (LibreOffice → PDF → pdftoppm);
- ``artifacts/voiceover/slide-NN.mp3`` and ``artifacts/voiceover/architecture-prelude.mp4`` (1920x1080, H.264, AAC);
- ``artifacts/voiceover/build-summary.md`` — durations, the voice used, the visuals per slide, every TODO(verify).

Icons are Tabler outline SVGs (MIT), fetched once into ``docs/Decks/architecture-prelude/icons/``, recoloured to
the palette and rasterised at 512 px with cairosvg; without them a slide falls back to a drawn circle.
Voice, first that works: Amazon Polly neural, Polly standard (both in ``AWS_REGION``), edge-tts, pyttsx3.
Credentials come from the root ``.env`` (non-empty ``KEY=value`` lines, as ``mk/vars.mk`` loads them); nothing
is printed. Audio is cached per slide on a hash of voice + text.

Usage: ``uv run python scripts/build_prelude_deck.py [--no-video] [--skip-tts]``.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import sys
import urllib.request
import xml.etree.ElementTree as ET
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DECK_DIR = ROOT / "docs" / "Decks" / "architecture-prelude"
SCRIPT_MD = DECK_DIR / "voiceover-script.md"
DECK = DECK_DIR / "ask-the-tower-architecture-prelude.pptx"
ICON_DIR = DECK_DIR / "icons"
PREVIEW = DECK_DIR / "preview"
VOICE_DIR = ROOT / "artifacts" / "voiceover"
BUILD = VOICE_DIR / "build"
VIDEO = VOICE_DIR / "architecture-prelude.mp4"
SUMMARY = VOICE_DIR / "build-summary.md"
DIAGRAMS = ROOT / "docs" / "architecture" / "diagrams"

TABLER = "https://raw.githubusercontent.com/tabler/tabler-icons/main"
CARLITO = "https://raw.githubusercontent.com/google/fonts/main/ofl/carlito"

# Palette: accent = Tower / our code, grey = AWS managed, warn = attacker / risk, ok = evidence.
BG, PANEL, LINE = "0B1220", "141F38", "26365C"
INK, MUTED = "E6ECFF", "9AA8CC"
ACCENT, WARN, OK, GREY = "4CC9F0", "F4A261", "2EC27E", "8A94A6"
TONES = {"ours": ACCENT, "aws": GREY, "risk": WARN, "ok": OK}
W_IN, H_IN = 13.333, 7.5

# Brief limits: <= 5 items, <= 8 words a label (decision fields and captions <= 12), 90-130 spoken words a slide.
MAX_ITEMS, MAX_LABEL_WORDS, MAX_SENTENCE_WORDS, NARRATION_WORDS = 5, 8, 12, (90, 130)
SENTENCE_FIELDS = (
    "subtitle",
    "kicker",
    "caption",
    "note",
    "decision",
    "why",
    "rejected",
    "tradeoff",
    "strip",
)
SPEECH_RATE = "118%"
PAD_S = 0.5

# What must never reach a slide, a note or the redacted diagram (the deck's privacy rule).
FORBIDDEN = {
    "phone number": re.compile(r"\+?\d{10,15}"),
    "phone prefix": re.compile(r"\+\d"),
    "account id": re.compile(r"\b\d{12}\b|\d{4}…\d{4}"),
    "ARN": re.compile(r"\barn:aws"),
    "IP address": re.compile(r"\b\d{1,3}(?:\.\d{1,3}){3}\b"),
    "bearer token": re.compile(r"Bearer\s+\S{8,}"),
}

# The as-built diagram labels carry masked account ids, the VPC CIDR, the SMS sandbox number's prefix and the
# registrar; the deck's copy drops them (and the title/legend/footnote cells), the source diagram is untouched.
DIAGRAM_REDACTIONS = [
    (re.compile(r",?\s*account \d{4}…\d{4}"), ""),
    (re.compile(r"\d{1,3}(?:\.\d{1,3}){3}/\d+ · "), ""),
    (re.compile(r"\s*\(\+\d[^)]*\)"), ""),
    (re.compile(r"\s*\(domain at [^)]*\)"), ""),
]
DIAGRAM_RESIDUE = re.compile(r"\d{4}…\d{4}|\+\d|\b\d{1,3}(?:\.\d{1,3}){3}\b|domain at ")
DIAGRAM_DROP = {"title", "subtitle", "note"}
# Recolour by role, in tones that read on the diagram's white page.
OURS = {"tower", "agent", "proxy", "binding", "alerts", "reconcile"}
CARRIER = {"mock"}
PEOPLE = {"phone", "laptop", "ops"}
ROLE_STYLE = {
    "ours": ("#E3F6FD", "#1E9BD7", 2),
    "carrier": ("#FDEBDD", "#E07A2E", 2),
    "aws": ("#EEF0F3", "#8A94A6", 1),
    "people": ("#FFFFFF", "#B5BDCB", 1),
    "zone": ("#F7F8FA", "#C8CED8", 1),
}
EDGE_RECOLOUR = {"#2BB3A0": "#1E9BD7", "#E8A33D": "#1E9BD7"}


def rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


def run(argv: list[str], cwd: Path | None = None) -> str:
    return subprocess.run(argv, check=True, capture_output=True, text=True, cwd=cwd).stdout  # noqa: S603 — fixed tool argv, no shell


def detect_font() -> str:
    """Inter when installed, else Calibri (the brief's order)."""
    try:
        families = run(["fc-list", ":", "family"])
    except (OSError, subprocess.CalledProcessError):
        return "Calibri"
    return "Inter" if re.search(r"^Inter\b", families, re.M) else "Calibri"


FONT = detect_font()
MONO = "Consolas"


@dataclass
class SlideSpec:
    number: int
    title: str
    fields: dict[str, str] = field(default_factory=dict)
    items: list[list[str]] = field(default_factory=list)
    narration: str = ""
    verify: list[str] = field(default_factory=list)

    @property
    def layout(self) -> str:
        return self.fields.get("layout", "title")


# --- parsing and checks -----------------------------------------------------------------------------------------


def parse(md: str) -> list[SlideSpec]:
    """Read the slides from the script markdown (everything before the first `## NN ·` is the preamble)."""
    slides: list[SlideSpec] = []
    section = ""
    for raw in md.splitlines():
        line = raw.rstrip()
        head = re.match(r"^## (\d{2}) · (.+)$", line)
        if head:
            slides.append(SlideSpec(int(head.group(1)), head.group(2).strip()))
            section = "fields"
            continue
        if not slides:
            continue
        s = slides[-1]
        if line == "On screen:":
            section = "items"
        elif line.startswith("Narration"):
            section = "narration"
        elif line == "Verify:":
            section = "verify"
        elif section == "fields" and re.match(r"^[a-z0-9]+: ", line):
            key, value = line.split(": ", 1)
            s.fields[key] = value.strip()
        elif section == "items" and line.startswith("- "):
            s.items.append([part.strip() for part in line[2:].split("|")])
        elif section == "narration" and line:
            s.narration = f"{s.narration} {line}".strip()
        elif section == "verify" and line.startswith("- "):
            s.verify.append(line[2:].strip())
    return slides


def words(t: str) -> int:
    return len(re.findall(r"[\w'’.,%‹›/-]+", t))


def label_parts(s: SlideSpec) -> list[str]:
    """The on-screen labels the word limit applies to (icon names, numbers and verbatim speech excluded)."""
    if s.layout == "decision" and s.fields.get("visual") == "table":
        return [row[0] for row in s.items]
    return [part for row in s.items for part in row[1:]]


def check(slides: list[SlideSpec]) -> list[str]:
    """Problems that stop the build: the brief's limits and the privacy rule."""
    problems: list[str] = []
    for s in slides:
        tag = f"slide {s.number:02d}"
        if s.layout not in DRAW:
            problems.append(f"{tag}: unknown layout {s.layout!r}")
        if len(s.items) > MAX_ITEMS:
            problems.append(f"{tag}: {len(s.items)} items (max {MAX_ITEMS})")
        for part in label_parts(s):
            if words(part) > MAX_LABEL_WORDS:
                problems.append(f"{tag}: label has {words(part)} words (max {MAX_LABEL_WORDS}): {part}")
        for key in SENTENCE_FIELDS:
            if key in s.fields and words(s.fields[key]) > MAX_SENTENCE_WORDS:
                problems.append(f"{tag}: {key} has {words(s.fields[key])} words (max {MAX_SENTENCE_WORDS})")
        n = words(s.narration)
        if not NARRATION_WORDS[0] <= n <= NARRATION_WORDS[1]:
            problems.append(
                f"{tag}: narration has {n} words (want {NARRATION_WORDS[0]}–{NARRATION_WORDS[1]})"
            )
        texts = [s.title, *s.fields.values(), *(p for row in s.items for p in row), s.narration, *s.verify]
        for t in texts:
            for what, rx in FORBIDDEN.items():
                if rx.search(t):
                    problems.append(f"{tag}: looks like a {what}: {t[:60]}…")
    return problems


# --- icons ------------------------------------------------------------------------------------------------------


def fetch(url: str, dst: Path) -> bool:
    """Download once; a present file is never fetched again. False when the network refuses."""
    if dst.exists() and dst.stat().st_size > 0:
        return True
    dst.parent.mkdir(parents=True, exist_ok=True)
    try:
        with urllib.request.urlopen(url, timeout=20) as r:  # noqa: S310 — fixed https URL
            dst.write_bytes(r.read())
        return True
    except OSError as e:
        print(
            f"prelude: could not fetch {url.rsplit('/', 1)[-1]} ({type(e).__name__}); drawing a shape instead"
        )
        return False


def icon_names(slides: list[SlideSpec]) -> set[str]:
    names = {
        "device-speaker",
        "device-mobile",
        "antenna-bars-5",
        "lock",
        "circle-x",
        "circle-check",
        "user",
        "player-play",
        "alert-triangle",
        "message-circle",
        "stopwatch",
        "route",
    }
    names |= {s.fields["icon"] for s in slides if "icon" in s.fields}
    for s in slides:
        if s.layout in {"flow", "threeup", "cards"} or s.fields.get("visual") in {
            "flow",
            "beforeafter",
            "datapath",
        }:
            names |= {row[0] for row in s.items}
    return names


def fetch_icons(names: set[str]) -> None:
    fetch(f"{TABLER}/LICENSE", ICON_DIR / "LICENSE-tabler-icons.txt")
    for name in sorted(names):
        fetch(f"{TABLER}/icons/outline/{name}.svg", ICON_DIR / f"{name}.svg")


def icon_png(name: str, color: str) -> Path | None:
    """The Tabler SVG recoloured and rasterised at 512 px (cached under the build folder)."""
    svg = ICON_DIR / f"{name}.svg"
    if not svg.exists():
        return None
    out = BUILD / "icons" / f"{name}-{color}.png"
    if not out.exists():
        import cairosvg

        out.parent.mkdir(parents=True, exist_ok=True)
        t = svg.read_text(encoding="utf-8").replace("currentColor", f"#{color}")
        cairosvg.svg2png(bytestring=t.encode(), write_to=str(out), output_width=512, output_height=512)
    return out


# --- the redacted, recoloured as-built diagram ------------------------------------------------------------------


def _role(cell_id: str) -> str:
    if cell_id in OURS:
        return "ours"
    if cell_id in CARRIER:
        return "carrier"
    if cell_id in PEOPLE:
        return "people"
    return "zone" if cell_id.startswith(("z_", "vpc_")) else "aws"


def _restyle(style: str, fill: str, stroke: str, width: int) -> str:
    style = re.sub(r"fillColor=#[0-9A-Fa-f]{6}", f"fillColor={fill}", style)
    style = re.sub(r"strokeColor=#[0-9A-Fa-f]{6}", f"strokeColor={stroke}", style)
    return re.sub(r"strokeWidth=\d+;?", "", style).rstrip(";") + f";strokeWidth={width};"


CALLOUT_STYLE = (
    "ellipse;whiteSpace=wrap;html=1;fillColor=#4CC9F0;strokeColor=#0B1220;strokeWidth=2;"
    "fontColor=#0B1220;fontStyle=1;fontSize=24;align=center;verticalAlign=middle;"
)


def prepare_diagram(stem: str, callouts: list[str]) -> str:
    """The deck's copy of the diagram: redacted, recoloured by role, titles dropped, and a numbered callout on the
    top-right corner of each box in `callouts` (drawn as diagram cells, so the renderer places them exactly)."""
    xml = (DIAGRAMS / f"{stem}.drawio").read_text(encoding="utf-8")
    for rx, repl in DIAGRAM_REDACTIONS:
        xml = rx.sub(repl, xml)
    root = ET.fromstring(xml)  # noqa: S314 — the repo's own generated diagram
    boxes: dict[str, tuple[float, float, float, float]] = {}
    for parent in list(root.iter()):
        for cell in list(parent):
            if cell.tag != "mxCell":
                continue
            cid, style, geo = cell.get("id", ""), cell.get("style", ""), cell.find("mxGeometry")
            if cid in DIAGRAM_DROP:
                parent.remove(cell)
            elif cell.get("vertex") == "1" and geo is not None:
                if "text;" not in style:
                    cell.set("style", _restyle(style, *ROLE_STYLE[_role(cid)]))
                boxes[cid] = tuple(float(geo.get(k, 0)) for k in ("x", "y", "width", "height"))  # type: ignore[assignment]
            elif cell.get("edge") == "1":
                for old, new in EDGE_RECOLOUR.items():
                    style = style.replace(old, new)
                cell.set("style", style)
    graph = next(root.iter("root"))
    for i, cid in enumerate(callouts, start=1):
        x, y, w, _ = boxes[cid]
        cell = ET.SubElement(
            graph, "mxCell", id=f"callout{i}", value=str(i), style=CALLOUT_STYLE, vertex="1", parent="1"
        )
        ET.SubElement(
            cell,
            "mxGeometry",
            x=str(x + w - 28),
            y=str(y - 22),
            width="46",
            height="46",
            **{"as": "geometry"},
        )
    out = ET.tostring(root, encoding="unicode")
    labels = " ".join(html.unescape(v) for v in re.findall(r'value="([^"]*)"', out))
    leftover = DIAGRAM_RESIDUE.search(labels)
    if leftover:
        sys.exit(f"prelude: redaction of {stem} left {leftover.group(0)!r}; extend DIAGRAM_REDACTIONS")
    return out


def render_diagram(stem: str, xml: str) -> Path:
    out_dir = BUILD / "diagrams"
    out_dir.mkdir(parents=True, exist_ok=True)
    copy_path, png = out_dir / f"{stem}.drawio", out_dir / f"{stem}.png"
    if png.exists() and copy_path.exists() and copy_path.read_text(encoding="utf-8") == xml:
        return png
    copy_path.write_text(xml, encoding="utf-8")
    sys.path.insert(0, str(ROOT / "scripts"))
    import diagrams_png

    diagrams_png.OUT = out_dir
    diagrams_png.export_browser([copy_path])
    return png


# --- drawing primitives -----------------------------------------------------------------------------------------


Visuals = list[str]


def text(
    slide: Slide,
    x: float,
    y: float,
    w: float,
    h: float,
    runs: list[tuple[str, int, str, bool]],
    *,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    font: str | None = None,
    space: int = 6,
) -> None:
    """One text box; each run is its own paragraph: (text, pt, colour, bold)."""
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    for i, (t, size, color, bold) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = t
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, font or FONT
        r.font.color.rgb = rgb(color)


def shape(
    slide: Slide,
    kind: MSO_SHAPE,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str | None = PANEL,
    line: str | None = LINE,
    width: float = 1.25,
    dash: bool = False,
    radius: float | None = 0.08,
):  # noqa: ANN201
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = radius
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(width)
        if dash:
            sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    sp.shadow.inherit = False
    return sp


def card(slide: Slide, x: float, y: float, w: float, h: float, edge: str = LINE) -> None:
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, line=edge)


def icon(slide: Slide, name: str, x: float, y: float, size: float, color: str, vis: Visuals) -> None:
    """A recoloured Tabler icon, or a drawn ring with the name's initial when the SVG is unavailable."""
    png = icon_png(name, color)
    if png is not None:
        slide.shapes.add_picture(str(png), Inches(x), Inches(y), Inches(size), Inches(size))
        vis.append(f"icon {name}")
        return
    shape(slide, MSO_SHAPE.OVAL, x, y, size, size, fill=None, line=color, width=2.5)
    text(
        slide,
        x,
        y,
        size,
        size,
        [(name[:1].upper(), int(size * 30), color, True)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    vis.append(f"shape ({name} unavailable)")


def arrow(slide: Slide, x: float, y: float, w: float, h: float, color: str) -> None:
    shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, w, h, fill=color, line=None)


def line(
    slide: Slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    color: str,
    *,
    width: float = 2.0,
    dash: bool = False,
) -> None:
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT


def pill(
    slide: Slide,
    x: float,
    y: float,
    w: float,
    h: float,
    label: str,
    color: str,
    vis: Visuals,
    icon_name: str | None = None,
    size: int = 18,
) -> None:
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, line=color, width=1.5, radius=0.5)
    tx = x + 0.3
    if icon_name:
        icon(slide, icon_name, x + 0.25, y + (h - 0.42) / 2, 0.42, color, vis)
        tx = x + 0.8
    text(slide, tx, y, w - (tx - x) - 0.25, h, [(label, size, color, False)], anchor=MSO_ANCHOR.MIDDLE)


def badge(slide: Slide, x: float, y: float, n: int, size: float = 0.42) -> None:
    shape(slide, MSO_SHAPE.OVAL, x, y, size, size, fill=ACCENT, line=BG, width=1.5)
    text(
        slide,
        x,
        y,
        size,
        size,
        [(str(n), int(size * 38), BG, True)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def header(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    x = 0.6
    if "icon" in s.fields:
        icon(slide, s.fields["icon"], 0.6, 0.42, 0.62, ACCENT, vis)
        x = 1.4
    text(slide, x, 0.3, W_IN - x - 0.6, 0.85, [(s.title, 32, INK, True)], anchor=MSO_ANCHOR.MIDDLE)


def chrome(slide: Slide, s: SlideSpec, total: int) -> None:
    """Footer strip (number • Ask the Tower • section) and the progress bar along the bottom edge."""
    line(slide, 0.6, 6.95, W_IN - 0.6, 6.95, LINE, width=0.75)
    footer = f"{s.number:02d}  •  Ask the Tower  •  {s.fields.get('section', '')}"
    text(slide, 0.6, 6.98, 9.0, 0.34, [(footer, 11, MUTED, False)], anchor=MSO_ANCHOR.MIDDLE)
    shape(slide, MSO_SHAPE.RECTANGLE, 0, H_IN - 0.07, W_IN, 0.07, fill=PANEL, line=None)
    shape(slide, MSO_SHAPE.RECTANGLE, 0, H_IN - 0.07, W_IN * s.number / total, 0.07, fill=ACCENT, line=None)


# --- layouts ----------------------------------------------------------------------------------------------------


def draw_title(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    icon(slide, "device-speaker", 1.3, 1.25, 1.7, ACCENT, vis)
    icon(slide, "device-mobile", 10.35, 1.25, 1.7, INK, vis)
    icon(slide, "antenna-bars-5", 11.75, 0.85, 0.75, OK, vis)
    line(slide, 3.2, 2.1, 10.15, 2.1, ACCENT, width=3, dash=True)
    vis.append("dotted link speaker ↔ phone")
    text(slide, 0.8, 3.15, W_IN - 1.6, 1.3, [(s.title, 66, INK, True)], align=PP_ALIGN.CENTER)
    text(slide, 1.2, 4.5, W_IN - 2.4, 0.8, [(s.fields["subtitle"], 26, MUTED, False)], align=PP_ALIGN.CENTER)
    text(slide, 1.2, 5.35, W_IN - 2.4, 0.5, [(s.fields["kicker"], 20, ACCENT, False)], align=PP_ALIGN.CENTER)
    if "badge" in s.fields:
        pill(
            slide, (W_IN - 3.6) / 2, 6.1, 3.6, 0.55, s.fields["badge"], MUTED, vis, "message-circle", size=16
        )


def draw_flow(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    header(slide, s, vis)
    w, gap = 4.25, -0.25
    for i, (name, label, *_) in enumerate(s.items):
        x = 0.6 + i * (w + gap)
        kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        sp = shape(slide, kind, x, 1.45, w, 2.55, fill=PANEL, line=WARN, width=2, radius=None)
        sp.adjustments[0] = 0.22
        left = 0.35 if i == 0 else 0.85
        inner = w - left - 0.85
        icon(slide, name, x + left + (inner - 0.95) / 2, 1.65, 0.95, WARN, vis)
        text(slide, x + left, 2.7, inner, 1.2, [(label, 20, INK, True)], align=PP_ALIGN.CENTER)
    vis.append(f"{len(s.items)}-step chevron flow (warning)")
    if "number" in s.fields:
        text(slide, 0.6, 4.3, 3.6, 1.6, [(s.fields["number"], 96, WARN, True)], anchor=MSO_ANCHOR.MIDDLE)
        text(slide, 4.2, 4.45, 8.5, 0.7, [(s.fields["caption"], 22, INK, False)], anchor=MSO_ANCHOR.MIDDLE)
        vis.append(f"stat {s.fields['number']}")
    if "note" in s.fields:
        text(slide, 4.2, 5.2, 8.5, 0.6, [(s.fields["note"], 20, MUTED, False)], anchor=MSO_ANCHOR.MIDDLE)


def draw_threeup(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    header(slide, s, vis)
    w, gap = 3.85, 0.29
    for i, (name, tool, answer) in enumerate(s.items):
        x = 0.6 + i * (w + gap)
        card(slide, x, 1.45, w, 3.85, edge=ACCENT)
        shape(slide, MSO_SHAPE.OVAL, x + (w - 1.35) / 2, 1.7, 1.35, 1.35, fill=BG, line=ACCENT, width=1.5)
        icon(slide, name, x + (w - 0.85) / 2, 1.95, 0.85, ACCENT, vis)
        text(slide, x + 0.2, 3.2, w - 0.4, 0.6, [(tool, 26, ACCENT, True)], align=PP_ALIGN.CENTER, font=MONO)
        text(slide, x + 0.25, 3.9, w - 0.5, 1.3, [(answer, 20, INK, False)], align=PP_ALIGN.CENTER)
    vis.append("three-up tool columns")
    pill(slide, 0.6, 5.65, W_IN - 1.2, 0.75, s.fields["badge"], ACCENT, vis, "lock", size=20)
    vis.append("consent-first lock badge")


def draw_picture(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    header(slide, s, vis)
    stem = s.fields["image"].split(":", 1)[1]
    xml = prepare_diagram(stem, [row[0] for row in s.items])
    png = render_diagram(stem, xml)
    with Image.open(png) as im:
        ratio = im.width / im.height
    max_w, max_h = 9.55, 5.45
    w, h = (max_w, max_w / ratio) if max_w / ratio <= max_h else (max_h * ratio, max_h)
    px, py = 0.55, 1.3 + (max_h - h) / 2
    shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        px - 0.08,
        py - 0.08,
        w + 0.16,
        h + 0.16,
        fill="FFFFFF",
        line=None,
        radius=0.03,
    )
    slide.shapes.add_picture(str(png), Inches(px), Inches(py), Inches(w), Inches(h))
    vis.append(f"{stem} diagram (redacted, recoloured by role)")
    lx = px + w + 0.3
    lw = W_IN - lx - 0.45
    for i, (_cell, label) in enumerate(s.items, start=1):
        y = 1.35 + (i - 1) * 1.2
        badge(slide, lx, y, i)
        text(slide, lx + 0.55, y - 0.08, lw - 0.55, 1.1, [(label, 18, INK, False)])
    vis.append(f"{len(s.items)} numbered callouts (in the diagram) + legend")
    text(slide, lx, 6.0, lw, 0.85, [(s.fields["caption"], 14, MUTED, False)])


def draw_decision(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    """Left: the decision and why. Right: the rejected alternative with ✗. Middle: the picture. Bottom: trade-off."""
    header(slide, s, vis)
    f = s.fields
    card(slide, 0.6, 1.3, 7.75, 1.6, edge=ACCENT)
    text(
        slide,
        0.85,
        1.38,
        7.3,
        1.45,
        [("DECISION", 13, ACCENT, True), (f["decision"], 22, INK, True), (f["why"], 18, MUTED, False)],
        space=3,
    )
    card(slide, 8.55, 1.3, 4.18, 1.6, edge=WARN)
    icon(slide, "circle-x", 8.75, 1.45, 0.42, WARN, vis)
    text(slide, 9.25, 1.42, 3.4, 0.4, [("REJECTED", 13, WARN, True)], anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 8.8, 1.9, 3.8, 0.95, [(f["rejected"], 18, INK, False)])
    VISUAL[f["visual"]](slide, s, vis)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 6.12, W_IN - 1.2, 0.62, fill=PANEL, line=None, radius=0.3)
    text(
        slide,
        0.85,
        6.12,
        W_IN - 1.7,
        0.62,
        [(f"TRADE-OFF   {f['tradeoff']}", 18, INK, False)],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    vis.append("decision card: decision/why · rejected ✗ · trade-off strip")


BAND_Y, BAND_H = 3.1, 2.85


def visual_flow(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    n = len(s.items)
    step, w = (W_IN - 1.2) / n, 2.3
    for i, (name, label, *tone) in enumerate(s.items):
        color = TONES.get(tone[0] if tone else "ours", ACCENT)
        cx = 0.6 + i * step + step / 2
        shape(slide, MSO_SHAPE.OVAL, cx - 0.6, BAND_Y + 0.2, 1.2, 1.2, fill=PANEL, line=color, width=2)
        icon(slide, name, cx - 0.36, BAND_Y + 0.44, 0.72, color, vis)
        text(slide, cx - w / 2, BAND_Y + 1.55, w, 1.2, [(label, 18, INK, False)], align=PP_ALIGN.CENTER)
        if i < n - 1:
            arrow(slide, cx + 0.8, BAND_Y + 0.62, step - 1.6, 0.36, MUTED)
    vis.append(f"{n}-step request flow")


def visual_beforeafter(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    (bad_icon, bad), (good_icon, good), (extra_icon, extra) = (row[:2] for row in s.items)
    for x, name, label, color, tag in (
        (0.6, bad_icon, bad, WARN, "BEFORE"),
        (5.0, good_icon, good, OK, "AFTER"),
    ):
        card(slide, x, BAND_Y + 0.15, 3.8, 2.5, edge=color)
        text(slide, x + 0.25, BAND_Y + 0.25, 3.3, 0.4, [(tag, 13, color, True)])
        icon(slide, name, x + 1.45, BAND_Y + 0.65, 0.9, color, vis)
        text(slide, x + 0.25, BAND_Y + 1.65, 3.3, 0.95, [(label, 20, INK, True)], align=PP_ALIGN.CENTER)
    arrow(slide, 4.48, BAND_Y + 1.15, 0.45, 0.42, MUTED)
    card(slide, 9.1, BAND_Y + 0.15, 3.63, 2.5, edge=GREY)
    icon(slide, extra_icon, 10.47, BAND_Y + 0.4, 0.9, GREY, vis)
    text(slide, 9.3, BAND_Y + 1.45, 3.25, 1.1, [(extra, 20, INK, False)], align=PP_ALIGN.CENTER)
    vis.append("before/after pair + CAMARA targets card")


def visual_datapath(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    tones = [WARN, ACCENT, GREY, OK]
    n = len(s.items)
    step = (W_IN - 1.2 + 0.4) / n
    for i, (name, label, *_) in enumerate(s.items):
        color = tones[i % len(tones)]
        x = 0.6 + i * step
        card(slide, x, BAND_Y, step - 0.4, 1.0, edge=color)
        icon(slide, name, x + 0.2, BAND_Y + 0.2, 0.6, color, vis)
        text(slide, x + 0.95, BAND_Y, step - 1.45, 1.0, [(label, 18, INK, True)], anchor=MSO_ANCHOR.MIDDLE)
        if i < n - 1:
            arrow(slide, x + step - 0.37, BAND_Y + 0.36, 0.34, 0.28, MUTED)
    vis.append("data path: number → HMAC → KMS → never logged")
    top = BAND_Y + 1.3
    for x, who in ((0.75, "Mom"), (5.35, "Asish")):
        icon(slide, "user", x + 0.2, top + 0.2, 0.6, INK, vis)
        text(slide, x - 0.1, top + 0.85, 1.2, 0.45, [(who, 18, INK, True)], align=PP_ALIGN.CENTER)
    arrow(slide, 1.95, top + 0.32, 3.2, 0.26, OK)
    text(slide, 1.95, top - 0.08, 3.2, 0.38, [("grants per tool", 16, OK, False)], align=PP_ALIGN.CENTER)
    shape(slide, MSO_SHAPE.LEFT_ARROW, 1.95, top + 0.78, 3.2, 0.26, fill=WARN, line=None, radius=None)
    text(slide, 1.95, top + 1.06, 3.2, 0.38, [("revokes any time", 16, WARN, False)], align=PP_ALIGN.CENTER)
    vis.append("grant/revoke arrows Mom ↔ Asish")
    x0 = 7.3
    text(slide, x0, top - 0.08, 5.4, 0.38, [("AUDIT, BEFORE EVERY ANSWER", 13, ACCENT, True)])
    for i in range(4):
        x = x0 + i * 1.12
        card(slide, x, top + 0.35, 0.8, 0.7, edge=ACCENT)
        text(
            slide,
            x,
            top + 0.35,
            0.8,
            0.7,
            [(f"#{i + 1}", 18, INK, True)],
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        if i < 3:
            line(slide, x + 0.8, top + 0.7, x + 1.12, top + 0.7, ACCENT, width=2.5)
    text(slide, x0 + 4.5, top + 0.35, 1.0, 0.7, [("hash-linked", 14, MUTED, False)], anchor=MSO_ANCHOR.MIDDLE)
    vis.append("audit hash chain as linked blocks")


def _cell(cell, value: str, size: int, color: str, bold: bool, fill: str, font: str | None = None) -> None:  # noqa: ANN001
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    cell.margin_left = cell.margin_right = Inches(0.12)
    cell.margin_top = cell.margin_bottom = Inches(0.04)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    tf.text = ""
    r = tf.paragraphs[0].add_run()
    r.text = value
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, font or FONT
    r.font.color.rgb = rgb(color)


def visual_table(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    tbl = slide.shapes.add_table(
        len(s.items) + 1, 2, Inches(0.6), Inches(BAND_Y - 0.05), Inches(W_IN - 1.2), Inches(BAND_H)
    ).table
    tbl.columns[0].width = Inches(3.3)
    tbl.columns[1].width = Inches(W_IN - 1.2 - 3.3)
    heights = [0.42] + [0.98 if len(r[1]) > 60 else 0.5 for r in s.items]
    for r, h in enumerate(heights):
        tbl.rows[r].height = Inches(h)
    for c, label in enumerate(("Reason code", "Spoken, word for word (phrasing.py)")):
        _cell(tbl.cell(0, c), label, 14, ACCENT, True, LINE)
    for r, (code, spoken, *_) in enumerate(s.items, start=1):
        _cell(tbl.cell(r, 0), code, 16, ACCENT, True, PANEL, font=MONO)
        _cell(tbl.cell(r, 1), f"“{spoken}”", 18, INK, False, PANEL)
    vis.append(f"decision-table excerpt ({len(s.items)} reason codes)")


def visual_budget(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    v = {row[0]: (int(row[1]), row[2]) for row in s.items}
    scale_max, x0, width = 600, 0.9, 11.4
    bar_y, bar_h = BAND_Y + 0.75, 0.6

    def at(ms: float) -> float:
        return x0 + width * ms / scale_max

    fixed, carrier = v["fixed"][0], v["carrier"][0]
    shape(slide, MSO_SHAPE.RECTANGLE, at(0), bar_y, at(fixed) - at(0), bar_h, fill=GREY, line=None)
    shape(
        slide,
        MSO_SHAPE.RECTANGLE,
        at(fixed),
        bar_y,
        at(fixed + carrier) - at(fixed),
        bar_h,
        fill=WARN,
        line=None,
    )
    text(
        slide,
        at(0),
        bar_y,
        at(fixed) - at(0),
        bar_h,
        [(f"≈ {fixed} ms", 18, BG, True)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    text(
        slide,
        at(fixed),
        bar_y,
        at(fixed + carrier) - at(fixed),
        bar_h,
        [(f"≈ {carrier}", 18, BG, True)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    text(slide, at(0), bar_y - 0.48, at(fixed) - at(0), 0.45, [(v["fixed"][1], 16, MUTED, False)])
    text(
        slide,
        at(fixed + carrier) - 4.6,
        bar_y + bar_h + 0.05,
        4.6,
        0.45,
        [(v["carrier"][1], 16, WARN, False)],
        align=PP_ALIGN.RIGHT,
    )
    target = v["target"][0]
    line(slide, at(target), bar_y - 0.6, at(target), bar_y + bar_h, INK, width=2, dash=True)
    text(
        slide,
        at(target) - 1.5,
        bar_y - 0.98,
        3.0,
        0.4,
        [(f"{v['target'][1]} {target} ms", 16, INK, True)],
        align=PP_ALIGN.CENTER,
    )
    total = v["total"][0]
    right = W_IN - 0.6 - (at(total) + 0.1)
    text(
        slide,
        at(total) + 0.1,
        bar_y - 0.2,
        right,
        0.55,
        [(f"{total} ms", 22, WARN, True)],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    text(slide, at(total) + 0.1, bar_y + 0.3, right, 0.6, [(v["total"][1], 14, MUTED, False)])
    for ms in range(0, scale_max + 1, 100):
        text(
            slide,
            at(ms) - 0.4,
            bar_y + bar_h + 0.48,
            0.8,
            0.3,
            [(str(ms), 12, MUTED, False)],
            align=PP_ALIGN.CENTER,
        )
    vis.append("latency budget bar with 400 ms target marker")
    stored = v["stored"]
    pill(slide, 0.6, BAND_Y + 2.15, 7.0, 0.62, f"{stored[1]}  ({stored[0]} ms carrier)", OK, vis, "stopwatch")


def draw_cards(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    header(slide, s, vis)
    n, gap = len(s.items), 0.25
    w = (W_IN - 1.2 - gap * (n - 1)) / n
    for i, (name, title, detail) in enumerate(s.items):
        x = 0.6 + i * (w + gap)
        big = "→" in title
        card(slide, x, 1.35, w, 4.35, edge=ACCENT if big else LINE)
        icon(slide, name, x + 0.3, 1.6, 0.8, ACCENT, vis)
        text(slide, x + 0.25, 2.6, w - 0.5, 1.4, [(title, 34 if big else 22, ACCENT if big else INK, True)])
        text(slide, x + 0.25, 4.0, w - 0.5, 1.6, [(detail, 18, MUTED, False)])
    vis.append(f"{n} lesson cards")
    if "strip" in s.fields:
        pill(slide, 0.6, 5.95, W_IN - 1.2, 0.68, s.fields["strip"], MUTED, vis, "route")


def draw_evidence(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    header(slide, s, vis)
    data = CategoryChartData()
    data.categories = [row[0] for row in s.items]
    data.add_series("Passed", [int(row[1]) for row in s.items])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.3), Inches(7.4), Inches(5.45), data
    ).chart
    chart.has_legend = False
    chart.has_title = True
    chart.chart_title.text_frame.text = "Tests passed, by layer (0 failed)"
    tr = chart.chart_title.text_frame.paragraphs[0].runs[0]
    tr.font.size, tr.font.bold, tr.font.name = Pt(18), True, FONT
    tr.font.color.rgb = rgb(INK)
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold, dl.font.name = Pt(18), True, FONT
    dl.font.color.rgb = rgb(INK)
    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = rgb(OK)
    cat = chart.category_axis
    cat.tick_labels.font.size, cat.tick_labels.font.name = Pt(16), FONT
    cat.tick_labels.font.color.rgb = rgb(INK)
    cat.format.line.color.rgb = rgb(LINE)
    val = chart.value_axis
    val.visible = False
    val.has_major_gridlines = False
    val.maximum_scale = max(int(row[1]) for row in s.items) * 1.15
    vis.append("column chart of tests passed by layer (native chart)")
    for i in (1, 2, 3):
        tone, label = (p.strip() for p in s.fields[f"badge{i}"].split("|", 1))
        color = OK if tone == "ok" else WARN
        y = 1.45 + (i - 1) * 1.75
        card(slide, 8.3, y, 4.43, 1.5, edge=color)
        icon(slide, "circle-check" if tone == "ok" else "alert-triangle", 8.5, y + 0.43, 0.65, color, vis)
        text(slide, 9.3, y, 3.3, 1.5, [(label, 20, INK, True)], anchor=MSO_ANCHOR.MIDDLE)
    vis.append("three evidence badges (green pass, amber over target)")


def draw_handover(slide: Slide, s: SlideSpec, vis: Visuals) -> None:
    header(slide, s, vis)
    shape(slide, MSO_SHAPE.CAN, 1.0, 1.75, 2.4, 3.5, fill="1C2A4A", line=LINE, radius=None).adjustments[0] = (
        0.18
    )
    shape(slide, MSO_SHAPE.OVAL, 1.0, 1.75, 2.4, 0.63, fill=None, line=ACCENT, width=4.5)
    icon(slide, "device-speaker", 1.75, 3.2, 0.9, MUTED, vis)
    text(slide, 0.6, 5.35, 3.2, 0.5, [("Echo, on home Wi-Fi", 18, MUTED, False)], align=PP_ALIGN.CENTER)
    vis.append("stylised Echo with light ring")
    for i, (question, answer) in enumerate(s.items):
        y = 1.4 + i * 1.42
        sp = shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT,
            4.3,
            y,
            8.43,
            1.18,
            fill=PANEL,
            line=ACCENT,
            radius=None,
        )
        sp.adjustments[0], sp.adjustments[1] = -0.56, 0.1
        text(
            slide,
            4.55,
            y + 0.05,
            8.0,
            1.1,
            [(question, 22, INK, True), (f"→ {answer}", 18, ACCENT, False)],
            anchor=MSO_ANCHOR.MIDDLE,
            space=2,
        )
    vis.append(f"{len(s.items)} speech bubbles (the Moments)")
    text(slide, 4.3, 5.75, 5.6, 0.7, [(s.fields["strip"], 20, MUTED, False)], anchor=MSO_ANCHOR.MIDDLE)
    shape(slide, MSO_SHAPE.OVAL, 10.0, 5.68, 0.82, 0.82, fill=ACCENT, line=None)
    icon(slide, "player-play", 10.16, 5.84, 0.5, BG, vis)
    text(slide, 10.95, 5.68, 1.9, 0.82, [(s.fields["cta"], 24, ACCENT, True)], anchor=MSO_ANCHOR.MIDDLE)


DRAW: dict[str, Callable[[Slide, SlideSpec, Visuals], None]] = {
    "title": draw_title,
    "flow": draw_flow,
    "threeup": draw_threeup,
    "picture": draw_picture,
    "decision": draw_decision,
    "cards": draw_cards,
    "evidence": draw_evidence,
    "handover": draw_handover,
}
VISUAL: dict[str, Callable[[Slide, SlideSpec, Visuals], None]] = {
    "flow": visual_flow,
    "beforeafter": visual_beforeafter,
    "datapath": visual_datapath,
    "table": visual_table,
    "budget": visual_budget,
}


# --- voice ------------------------------------------------------------------------------------------------------


def load_dotenv() -> None:
    """The repo's export loop: every non-empty KEY=value in the root .env goes into the environment."""
    env = ROOT / ".env"
    if not env.exists():
        return
    for raw in env.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^([A-Z0-9_]+)=(.+)$", raw.strip())
        if m:
            os.environ[m.group(1)] = m.group(2)


def _ssml(t: str) -> str:
    return f'<speak><prosody rate="{SPEECH_RATE}">{html.escape(t, quote=False)}</prosody></speak>'


def _polly(engine: str, voice: str) -> tuple[str, object]:
    import boto3

    region = os.environ.get("AWS_REGION", "us-east-2")
    client = boto3.client("polly", region_name=region)

    def say(t: str, dst: Path) -> None:
        r = client.synthesize_speech(
            Text=_ssml(t), TextType="ssml", VoiceId=voice, Engine=engine, OutputFormat="mp3"
        )
        dst.write_bytes(r["AudioStream"].read())

    return f"Amazon Polly {engine}, {voice}, {region}, rate {SPEECH_RATE}", say


def _edge(voice: str) -> tuple[str, object]:
    import asyncio

    import edge_tts

    def say(t: str, dst: Path) -> None:
        asyncio.run(edge_tts.Communicate(t, voice, rate="+18%").save(str(dst)))

    return f"edge-tts {voice}", say


def _pyttsx3() -> tuple[str, object]:
    import pyttsx3

    def say(t: str, dst: Path) -> None:
        wav = dst.with_suffix(".wav")
        engine = pyttsx3.init()
        engine.save_to_file(t, str(wav))
        engine.runAndWait()
        run(
            [
                "ffmpeg",
                "-y",
                "-loglevel",
                "error",
                "-i",
                str(wav),
                "-codec:a",
                "libmp3lame",
                "-q:a",
                "3",
                str(dst),
            ]
        )
        wav.unlink()

    return "pyttsx3 (offline)", say


def pick_voice(probe_dir: Path, voice: str) -> tuple[str, object]:
    """The first engine that synthesises a test sentence; the same one is then used for every slide."""
    candidates = [
        lambda: _polly("neural", voice),
        lambda: _polly("standard", voice),
        lambda: _edge("en-US-GuyNeural"),
        _pyttsx3,
    ]
    probe = probe_dir / "probe.mp3"
    for make in candidates:
        try:
            label, say = make()
            say("Ask the Tower.", probe)  # type: ignore[operator]
            if probe.stat().st_size > 0:
                probe.unlink()
                return label, say
        except Exception as e:  # noqa: BLE001 — each engine fails differently; the next one is tried
            print(f"prelude: voice unavailable ({type(e).__name__}); trying the next one")
    sys.exit("prelude: no text-to-speech engine worked (Polly, edge-tts, pyttsx3)")


TTS_CACHE = BUILD / "tts-cache.json"


def cached_voices() -> dict[str, dict[str, str]]:
    return json.loads(TTS_CACHE.read_text()) if TTS_CACHE.exists() else {}


def voice_label(seen: dict[str, dict[str, str]]) -> str:
    return ", ".join(sorted({v["voice"] for v in seen.values()}))


def synthesise(slides: list[SlideSpec], voice: str) -> str:
    """slide-NN.mp3 for every slide whose text or voice changed; returns the voice(s) used."""
    load_dotenv()
    seen = cached_voices()
    say = None
    for s in slides:
        dst = VOICE_DIR / f"slide-{s.number:02d}.mp3"
        key, digest = f"{s.number:02d}", _digest(s.narration, voice)
        if dst.exists() and seen.get(key, {}).get("digest") == digest:
            continue
        if say is None:
            label, say = pick_voice(BUILD, voice)
        say(s.narration, dst)  # type: ignore[operator]
        seen[key] = {"voice": label, "digest": digest}
        print(f"prelude: {dst.relative_to(ROOT)} ({label})")
    TTS_CACHE.write_text(json.dumps(seen, indent=2))
    return voice_label(seen)


def _digest(t: str, voice: str) -> str:
    return hashlib.sha256(f"{voice}|{SPEECH_RATE}|{t}".encode()).hexdigest()[:16]


def duration(path: Path) -> float:
    out = run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(path)])
    return float(out.strip())


# --- deck -------------------------------------------------------------------------------------------------------


def autoplay(slide: Slide, shape_id: int, ms: int) -> None:
    """Start the slide's audio on slide start and advance when it ends (PowerPoint's own timing XML)."""
    p = "http://schemas.openxmlformats.org/presentationml/2006/main"
    sld = slide._element
    for old in sld.findall(f"{{{p}}}timing") + sld.findall(f"{{{p}}}transition"):
        sld.remove(old)
    transition = etree.fromstring(f'<p:transition xmlns:p="{p}" spd="med" advClick="1" advTm="{ms}"/>')
    timing = etree.fromstring(f"""<p:timing xmlns:p="{p}"><p:tnLst><p:par>
<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>
<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect">
<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{ms}" fill="hold"/>
<p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:audio><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/>
</p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl></p:cMediaNode></p:audio>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>""")
    ext = sld.find(f"{{{p}}}extLst")
    for el in (transition, timing):
        if ext is not None:
            ext.addprevious(el)
        else:
            sld.append(el)


def notes_text(s: SlideSpec) -> str:
    extra = "\n".join(f"[not spoken] {v}" for v in s.verify)
    return f"{s.narration}\n\n{extra}".strip()


def build_deck(
    slides: list[SlideSpec], dst: Path, audio: dict[int, tuple[Path, float]] | None
) -> dict[int, Visuals]:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN)
    visuals: dict[int, Visuals] = {}
    poster = icon_png("device-speaker", ACCENT)
    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(BG)
        vis: Visuals = []
        DRAW[s.layout](slide, s, vis)
        chrome(slide, s, len(slides))
        visuals[s.number] = [*vis, "footer + progress bar"]
        slide.notes_slide.notes_text_frame.text = notes_text(s)
        if audio and s.number in audio:
            mp3, secs = audio[s.number]
            media = slide.shapes.add_movie(
                str(mp3),
                Inches(W_IN - 0.9),
                Inches(6.98),
                Inches(0.3),
                Inches(0.3),
                poster_frame_image=str(poster) if poster else None,
                mime_type="audio/mpeg",
            )
            autoplay(slide, media.shape_id, int((secs + PAD_S) * 1000))
    dst.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dst))
    return visuals


# --- previews and video -----------------------------------------------------------------------------------------


def ensure_render_font() -> None:
    """LibreOffice renders a missing Calibri as DejaVu Sans, which is far wider; Carlito (OFL) has Calibri's metrics,
    so the previews and the video wrap text as PowerPoint does. Installed once under ~/.local/share/fonts."""
    if FONT != "Calibri" or re.search(r"Carlito|Calibri", run(["fc-list", ":", "family"])):
        return
    fonts = Path.home() / ".local" / "share" / "fonts" / "carlito"
    got = [
        fetch(f"{CARLITO}/Carlito-{style}.ttf", fonts / f"Carlito-{style}.ttf")
        for style in ("Regular", "Bold")
    ]
    if all(got):
        run(["fc-cache", "-f", str(fonts)])
        print("prelude: installed Carlito (Calibri metrics) for the LibreOffice render")


def render_pngs(silent: Path) -> list[Path]:
    """Every slide to PNG (LibreOffice → PDF → pdftoppm), copied to preview/ as slide-NN.png."""
    ensure_render_font()
    out = BUILD / "render"
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out), str(silent)])
    run(["pdftoppm", "-png", "-r", "150", str(out / f"{silent.stem}.pdf"), str(out / "slide")])
    pngs = sorted(out.glob("slide-*.png"))
    if PREVIEW.exists():
        shutil.rmtree(PREVIEW)
    PREVIEW.mkdir(parents=True)
    for i, png in enumerate(pngs, start=1):
        shutil.copy(png, PREVIEW / f"slide-{i:02d}.png")
    return pngs


def render_video(pngs: list[Path], audio: dict[int, tuple[Path, float]]) -> None:
    seg_dir = BUILD / "segments"
    seg_dir.mkdir(parents=True, exist_ok=True)
    listing = []
    for i, png in enumerate(pngs, start=1):
        mp3, secs = audio[i]
        seg = seg_dir / f"seg-{i:02d}.mp4"
        run(
            [
                "ffmpeg",
                "-y",
                "-loglevel",
                "error",
                "-loop",
                "1",
                "-framerate",
                "30",
                "-i",
                str(png),
                "-i",
                str(mp3),
                "-vf",
                "scale=1920:1080:flags=lanczos,format=yuv420p",
                "-af",
                f"apad=pad_dur={PAD_S}",
                "-t",
                f"{secs + PAD_S:.3f}",
                "-c:v",
                "libx264",
                "-tune",
                "stillimage",
                "-preset",
                "medium",
                "-crf",
                "20",
                "-r",
                "30",
                "-c:a",
                "aac",
                "-b:a",
                "160k",
                "-ar",
                "48000",
                "-ac",
                "2",
                str(seg),
            ]
        )
        listing.append(f"file '{seg.name}'")
    (seg_dir / "list.txt").write_text("\n".join(listing) + "\n")
    run(
        [
            "ffmpeg",
            "-y",
            "-loglevel",
            "error",
            "-f",
            "concat",
            "-safe",
            "0",
            "-i",
            "list.txt",
            "-c",
            "copy",
            "-movflags",
            "+faststart",
            str(VIDEO),
        ],
        cwd=seg_dir,
    )


# --- summary ----------------------------------------------------------------------------------------------------


def write_summary(
    slides: list[SlideSpec],
    audio: dict[int, tuple[Path, float]],
    voice: str,
    embedded: bool,
    visuals: dict[int, Visuals],
) -> str:
    total = sum(secs + PAD_S for _, secs in audio.values())
    mins, secs = divmod(round(total), 60)
    rows = [
        f"| {s.number:02d} | {s.title} | {words(s.narration)} | {audio[s.number][1] + PAD_S:.1f} |"
        for s in slides
        if s.number in audio
    ]
    vis_rows = []
    for s in slides:
        counted: dict[str, int] = {}
        for v in visuals.get(s.number, []):
            counted[v] = counted.get(v, 0) + 1
        listed = "; ".join(f"{v} ×{n}" if n > 1 else v for v, n in counted.items())
        vis_rows.append(f"- **{s.number:02d}** ({s.layout}): {listed}")
    todos = [f"- slide {s.number:02d}: {v}" for s in slides for v in s.verify if "TODO(verify)" in v]
    previews = sorted(PREVIEW.glob("*.png")) if PREVIEW.exists() else []
    paths = [DECK, *previews, *(a for a, _ in audio.values()), VIDEO if VIDEO.exists() else None]
    lines = [
        "# Architecture prelude — build summary",
        "",
        f"Generated by `scripts/build_prelude_deck.py` from `{SCRIPT_MD.relative_to(ROOT)}`.",
        "",
        f"- **Slides:** {len(slides)}",
        f"- **Total duration:** {total:.1f} s ({mins} min {secs:02d} s, {PAD_S} s padding a slide)",
        f"- **Voice:** {voice or 'not synthesised'}",
        f"- **Audio in the deck:** {'embedded, plays on slide start, auto-advance' if embedded else 'not embedded'}",
        f"- **Font:** {FONT} (rendered with Carlito where Calibri is missing)",
        "",
        "| # | Slide | Words | Seconds |",
        "|---|---|---:|---:|",
        *rows,
        "",
        "## Visuals per slide",
        "",
        *vis_rows,
        "",
        "## Files",
        "",
        *(f"- `{p.relative_to(ROOT)}`" for p in paths if p is not None),
        "",
        "## TODO(verify) left in the notes",
        "",
        *(todos or ["- none"]),
        "",
    ]
    out = "\n".join(lines)
    SUMMARY.write_text(out, encoding="utf-8")
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--skip-tts", action="store_true", help="use the MP3s already in artifacts/voiceover/")
    ap.add_argument(
        "--no-video",
        "--skip-video",
        dest="no_video",
        action="store_true",
        help="build the deck and the previews only",
    )
    ap.add_argument("--voice", default="Matthew", help="Polly voice id (default Matthew)")
    args = ap.parse_args()

    slides = parse(SCRIPT_MD.read_text(encoding="utf-8"))
    problems = check(slides)
    if problems:
        print("prelude: the script breaks the deck's rules:\n  " + "\n  ".join(problems))
        return 1
    BUILD.mkdir(parents=True, exist_ok=True)
    fetch_icons(icon_names(slides))

    voice = voice_label(cached_voices()) if args.skip_tts else synthesise(slides, args.voice)
    mp3s = {s.number: VOICE_DIR / f"slide-{s.number:02d}.mp3" for s in slides}
    audio = {n: (mp3, duration(mp3)) for n, mp3 in mp3s.items() if mp3.exists()}

    silent = BUILD / "prelude-silent.pptx"
    visuals = build_deck(slides, silent, None)
    embedded = len(audio) == len(slides)
    build_deck(slides, DECK, audio if embedded else None)
    print(
        f"prelude: {DECK.relative_to(ROOT)} ({len(slides)} slides, audio {'embedded' if embedded else 'absent'})"
    )
    pngs = render_pngs(silent)
    print(f"prelude: {len(pngs)} previews in {PREVIEW.relative_to(ROOT)}/")

    if not args.no_video:
        if not embedded:
            print("prelude: no video: some slides have no MP3 (run without --skip-tts)")
        else:
            render_video(pngs, audio)
            print(f"prelude: {VIDEO.relative_to(ROOT)} ({duration(VIDEO):.1f} s)")

    print(write_summary(slides, audio, voice, embedded, visuals))
    return 0


if __name__ == "__main__":
    sys.exit(main())
